"""Bounded server-side document rendering; no model-provided paths or HTML."""
from io import BytesIO
from pathlib import Path
import os


MIME = {
    "md": "text/markdown; charset=utf-8",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
}


def render_document(title: str, body: str, file_type: str) -> bytes:
    if file_type not in MIME:
        raise ValueError("Unsupported document format")
    title, body = title.strip()[:120], body.strip()[:100000]
    if not title or not body:
        raise ValueError("Document title and content are required")
    if file_type == "md":
        return f"# {title}\n\n{body}\n".encode("utf-8")
    output = BytesIO()
    if file_type == "docx":
        from docx import Document
        doc = Document()
        doc.add_heading(title, 0)
        for line in body.splitlines():
            doc.add_paragraph(line)
        doc.save(output)
        return output.getvalue()
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen.canvas import Canvas
    font_candidates = [os.environ.get("DOCUMENT_FONT_PATH", ""), "C:/Windows/Fonts/simhei.ttf", "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc"]
    font = next((Path(path) for path in font_candidates if path and Path(path).is_file()), None)
    if font is None:
        raise RuntimeError("PDF Chinese font is not configured")
    pdfmetrics.registerFont(TTFont("TongpinCJK", str(font)))
    canvas = Canvas(output, pagesize=A4)
    width, height = A4
    margin, y = 48, height - 55

    def write_line(text: str, size: int = 11):
        nonlocal y
        canvas.setFont("TongpinCJK", size)
        chars = max(1, int((width - 2 * margin) / (size * 1.05)))
        for start in range(0, max(1, len(text)), chars):
            if y < 50:
                canvas.showPage()
                y = height - 55
                canvas.setFont("TongpinCJK", size)
            canvas.drawString(margin, y, text[start:start + chars])
            y -= size * 1.55

    write_line(title, 17)
    y -= 14
    for line in body.splitlines():
        write_line(line)
    canvas.save()
    return output.getvalue()
